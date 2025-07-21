import streamlit as st
import requests
import docx
import PyPDF2
import nltk
from datetime import datetime
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from sentence_transformers import SentenceTransformer, util
from fpdf import FPDF

# Download required NLTK data
try:
    nltk.data.find('tokenizers/punkt')
    nltk.data.find('tokenizers/punkt_tab')
except LookupError:
    nltk.download('punkt')
    nltk.download('punkt_tab')

from nltk.tokenize import sent_tokenize

# Load BERT model
bert_model = SentenceTransformer('all-MiniLM-L6-v2')

# Your SerpAPI key
SERPAPI_KEY = "Put your own api key"

def extract_text_from_docx(file):
    doc = docx.Document(file)
    return '\n'.join([para.text for para in doc.paragraphs])

def extract_text_from_pdf(file):
    pdf = PyPDF2.PdfReader(file)
    return '\n'.join([page.extract_text() or '' for page in pdf.pages])

def get_serpapi_results(query, num_results=5):
    params = {
        "engine": "google",
        "q": query,
        "api_key": SERPAPI_KEY,
        "num": num_results
    }
    response = requests.get("https://serpapi.com/search", params=params)
    data = response.json()
    unique_sources = {}
    for result in data.get("organic_results", []):
        snippet = result.get("snippet", "")
        url = result.get("link", "")
        if snippet and url not in unique_sources:
            unique_sources[url] = {
                "snippet": snippet,
                "url": url,
                "title": result.get("title", "")
            }
    return list(unique_sources.values())    

def compute_cosine_similarity(input_text, sources):
    snippets = [src['snippet'] for src in sources]
    tfidf = TfidfVectorizer().fit_transform([input_text] + snippets)
    return cosine_similarity(tfidf[0:1], tfidf[1:]).flatten()

def compute_semantic_similarity(input_text, sources):
    snippets = [src['snippet'] for src in sources]
    embeddings = bert_model.encode([input_text] + snippets, convert_to_tensor=True)
    return util.cos_sim(embeddings[0], embeddings[1:]).cpu().numpy().flatten()

def calculate_plagiarism_score(cos_scores, sem_scores):
    combined = 0.4 * cos_scores + 0.6 * sem_scores
    return min(100, max(0, (combined.mean() * 100) * 1.25))

def get_top_matches(input_text, sources, model, top_n=10):
    input_sents = sent_tokenize(input_text)
    all_matches = []
    
    for src in sources:
        src_sents = sent_tokenize(src['snippet'])
        if not src_sents:
            continue
            
        embeddings = model.encode(input_sents + src_sents, convert_to_tensor=True)
        input_embeds = embeddings[:len(input_sents)]
        src_embeds = embeddings[len(input_sents):]
        
        sim_matrix = util.cos_sim(input_embeds, src_embeds)
        max_sim = sim_matrix.max(dim=1)
        
        for i in range(len(input_sents)):
            score = max_sim.values[i].item()
            if score > 0.4:
                all_matches.append({
                    'input': input_sents[i],
                    'source_sent': src_sents[max_sim.indices[i].item()],
                    'score': score,
                    'source_url': src['url'],
                    'source_title': src['title']
                })
    
    sorted_matches = sorted(all_matches, key=lambda x: x['score'], reverse=True)
    seen = set()
    unique_matches = []
    
    for match in sorted_matches:
        if match['source_url'] not in seen:
            seen.add(match['source_url'])
            unique_matches.append(match)
        if len(unique_matches) >= top_n:
            break
    
    return unique_matches[:top_n]

def generate_pdf_report(input_text, plagiarism_score, matches):
    pdf = FPDF()
    pdf.add_page()
    
    # Use built-in font without external files
    pdf.set_font("Arial", size=12)
    
    # Header
    pdf.cell(200, 10, txt="Plagiarism Check Report", ln=1, align='C')
    pdf.cell(200, 10, txt=f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=2)
    pdf.ln(10)
    
    # Summary
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(200, 10, txt="Summary", ln=1)
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt=f"Overall Plagiarism Score: {plagiarism_score:.1f}%", ln=1)
    pdf.ln(10)
    
    # Matches
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(200, 10, txt="Top Matches Found", ln=1)
    pdf.set_font("Arial", size=12)
    
    for idx, match in enumerate(matches, 1):
        # Clean text for PDF encoding
        def clean_text(text):
            return text.encode('latin-1', 'replace').decode('latin-1')
        
        pdf.multi_cell(0, 10, txt=clean_text(f"Match #{idx} (Similarity: {match['score']:.2f})"))
        pdf.multi_cell(0, 10, txt=clean_text(f"Source: {match['source_title']}"))
        pdf.multi_cell(0, 10, txt=clean_text(f"URL: {match['source_url']}"))
        pdf.multi_cell(0, 10, txt=clean_text("Your Content:"))
        pdf.multi_cell(0, 10, txt=clean_text(match['input']))
        pdf.multi_cell(0, 10, txt=clean_text("Matched Content:"))
        pdf.multi_cell(0, 10, txt=clean_text(match['source_sent']))
        pdf.ln(5)
    
    return pdf.output(dest='S').encode('latin-1', 'replace')

# Streamlit UI
st.set_page_config(page_title="Plagiarism Checker", page_icon="🧾")
st.title("🧾 Advanced Plagiarism Checker")

uploaded_file = st.file_uploader("📂 Upload a DOCX or PDF file", type=["docx", "pdf"])
input_text = ""

if uploaded_file:
    if uploaded_file.type == "application/pdf":
        input_text = extract_text_from_pdf(uploaded_file)
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        input_text = extract_text_from_docx(uploaded_file)

if not uploaded_file:
    input_text = st.text_area("✍️ Or paste text here:")

if st.button("🔍 Check for Plagiarism"):
    if input_text.strip():
        with st.spinner("Analyzing for potential matches..."):
            sources = get_serpapi_results(input_text[:300])
            
            if not sources:
                st.warning("No web results found. Try more content.")
            else:
                cos_scores = compute_cosine_similarity(input_text, sources)
                sem_scores = compute_semantic_similarity(input_text, sources)
                plagiarism_score = calculate_plagiarism_score(cos_scores, sem_scores)

                st.subheader("📊 Plagiarism Analysis")
                st.write(f"**Overall Plagiarism Score**: {plagiarism_score:.1f}%")
                
                if plagiarism_score > 70:
                    st.error("❗ High probability of plagiarism detected!")
                elif plagiarism_score > 40:
                    st.warning("⚠️ Significant similarities found")
                else:
                    st.success("✅ Content appears mostly original")

                top_matches = get_top_matches(input_text, sources, bert_model)
                
                if top_matches:
                    st.subheader("🔝 Top Unique Source Matches")
                    
                    for idx, match in enumerate(top_matches, 1):
                        st.markdown(f"**Match #{idx}** (Similarity: {match['score']:.2f})")
                        st.markdown(f"**Source**: [{match['source_title']}]({match['source_url']})")
                        
                        col1, col2 = st.columns(2)
                        with col1:
                            st.markdown("**Your Document:**")
                            st.warning(match['input'])
                        with col2:
                            st.markdown("**Web Content:**")
                            st.info(match['source_sent'])
                        
                        st.divider()

                    # Download Report Section
                    st.markdown("---")
                    st.subheader("📥 Download Report")
                    report = generate_pdf_report(input_text, plagiarism_score, top_matches)
                    
                    st.download_button(
                        label="Download PDF Report",
                        data=report,
                        file_name=f"plagiarism_report_{datetime.now().strftime('%Y%m%d%H%M%S')}.pdf",
                        mime="application/pdf"
                    )
                else:
                    st.info("No significant matches found")
    else:
        st.warning("Please upload a file or paste some text.")
